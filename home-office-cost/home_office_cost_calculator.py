#!/usr/bin/env python3
# -*- coding: utf-8 -*-

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

def create_home_office_cost_calculator():
    # 创建新的工作簿
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "成本计算器"

    # 设置列宽
    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['C'].width = 40
    ws.column_dimensions['D'].width = 15
    ws.column_dimensions['E'].width = 30

    # 设置标题样式
    title_font = Font(name='微软雅黑', size=12, bold=True)
    header_fill = PatternFill(start_color='CCE5FF', end_color='CCE5FF', fill_type='solid')

    # 添加标题
    ws['A1'] = '费用类型'
    ws['B1'] = '每月费用 (元)'
    ws['C1'] = '计算公式'
    ws['D1'] = '单项目分摊费用 (元)'
    ws['E1'] = '说明'

    # 设置标题格式
    for cell in ws[1]:
        cell.font = title_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')

    # 添加设备折旧计算
    ws['A2'] = '1. 笔记本电脑折旧'
    ws['A3'] = '原价'
    ws['B3'] = 15000
    ws['A4'] = '预计使用月数'
    ws['B4'] = 60
    ws['A5'] = '业务使用比例'
    ws['B5'] = 0.5
    ws['A6'] = '月折旧'
    ws['B6'] = '=B3/B4*B5'

    ws['A7'] = '2. 手机折旧'
    ws['A8'] = '原价'
    ws['B8'] = 5000
    ws['A9'] = '预计使用月数'
    ws['B9'] = 60
    ws['A10'] = '业务使用比例'
    ws['B10'] = 0.5
    ws['A11'] = '月折旧'
    ws['B11'] = '=B8/B9*B10'

    # 添加软件订阅费用
    ws['A13'] = '3. 软件订阅费用'
    ws['A14'] = '软件1（如视频剪辑会员）'
    ws['B14'] = 100
    ws['A15'] = '软件2（如 AI 助手订阅）'
    ws['B15'] = 200
    ws['A16'] = '月软件费用合计'
    ws['B16'] = '=SUM(B14:B15)'

    # 添加通信与网络费用
    ws['A18'] = '4. 通信与网络费用'
    ws['A19'] = '宽带月费'
    ws['B19'] = 200

    # 添加行政管理费用
    ws['A21'] = '5. 行政管理费用'
    ws['A22'] = '月度预算'
    ws['B22'] = 50

    # 添加总计
    ws['A24'] = '月度总成本'
    ws['B24'] = '=B6+B11+B16+B19+B22'

    # 添加项目分摊
    ws['A26'] = '每月项目数量'
    ws['B26'] = 10
    ws['A27'] = '单个项目分摊成本'
    ws['B27'] = '=B24/B26'

    # 保存文件
    wb.save('home_office_cost_calculator.xlsx')

if __name__ == '__main__':
    create_home_office_cost_calculator()
